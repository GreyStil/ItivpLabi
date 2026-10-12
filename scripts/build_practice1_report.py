# -*- coding: utf-8 -*-
"""Практическая работа №1 — отчёт (оформление как Lab1 WebRes)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(r"C:\Users\ASUS\Desktop\курс4\Итивп\Практическая1_Рубацкий.docx")

FONT = "Times New Roman"
SIZE = Pt(14)
LINE = Pt(18)
FIRST_LINE = Cm(1.25)


def set_run(run, *, bold=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = SIZE
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def apply_page_setup(section):
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)


def fmt_paragraph(p, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True):
    p.alignment = align
    p.paragraph_format.line_spacing = LINE
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    if indent and align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        p.paragraph_format.first_line_indent = FIRST_LINE
    else:
        p.paragraph_format.first_line_indent = Cm(0)


def add_paragraph(doc, text="", *, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True):
    p = doc.add_paragraph()
    fmt_paragraph(p, align=align, indent=indent)
    if text:
        set_run(p.add_run(text), bold=bold)
    return p


def add_body(doc, text):
    add_paragraph(doc, text)


def add_blank_line(doc):
    add_paragraph(doc, "", indent=False)


def add_section(doc, number: int, title: str):
    doc.add_page_break()
    add_paragraph(doc, f"{number} {title.upper()}", bold=True)
    add_paragraph(doc, "")


def add_subsection(doc, number: str, title: str):
    add_paragraph(doc, f"{number} {title}", bold=True, indent=True)


def add_field(paragraph, instruction: str):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def build_title_page(doc: Document):
    for text in (
        "Министерство образования Республики Беларусь",
        "",
        "Учреждение образования",
        "БЕЛОРУССКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ",
        "ИНФОРМАТИКИ И РАДИОЭЛЕКТРОНИКИ",
    ):
        add_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)

    add_paragraph(doc, "Факультет информационных технологий и управления", indent=True)
    add_paragraph(doc, "Кафедра информационных технологий автоматизированных систем", indent=True)
    add_paragraph(
        doc,
        "Дисциплина:  Интернет-технологии и веб-программирование",
        indent=True,
    )

    add_paragraph(doc, "ОТЧЁТ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    add_paragraph(doc, "к практической работе №1", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    add_paragraph(
        doc,
        "Серверное web-приложение на Express и EJS",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        indent=False,
    )
    add_paragraph(
        doc,
        "Тема курсового проекта: система анализа эффективности рекламных кампаний",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        indent=False,
    )

    add_paragraph(doc, "Выполнил", indent=True)
    add_paragraph(doc, "Р. С. Рубацкий", indent=True)
    add_paragraph(doc, "Проверила", indent=True)
    add_paragraph(doc, "Н. В. Хаджинова", indent=True)
    add_paragraph(doc, "Минск 2026", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)


def build_document():
    doc = Document()
    apply_page_setup(doc.sections[0])
    build_title_page(doc)

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_paragraph(doc, "СОДЕРЖАНИЕ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    p = doc.add_paragraph()
    fmt_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False)
    add_field(p, r'TOC \o "1-3" \h \z \u')
    add_body(
        doc,
        "Примечание: после открытия документа в Word выполните обновление полей содержания "
        "(правой кнопкой по содержанию — «Обновить поле»), чтобы номера страниц совпали с текстом.",
    )

    add_section(doc, 1, "Цель работы")
    add_body(
        doc,
        "Цель практической работы — освоить создание серверного web-приложения на Node.js "
        "и Express с шаблонизатором EJS, реализовать middleware для логирования и авторизации, "
        "организовать вывод динамических страниц на основе данных курсового проекта "
        "«Аналитика рекламных кампаний» с хранением данных в памяти сервера.",
    )

    add_section(doc, 2, "Ход выполнения работы")

    add_subsection(doc, "2.1", "Настройка проекта и EJS")
    add_body(
        doc,
        "Использован существующий проект курсовой работы (REST API рекламных кампаний). "
        "Установлены пакеты ejs, cookie-parser, bcrypt, jsonwebtoken, dotenv. "
        "В server.js настроены view engine и каталог шаблонов:",
    )
    add_paragraph(doc, "app.set('view engine', 'ejs');", indent=True)
    add_paragraph(doc, "app.set('views', './views');", indent=True)
    add_body(
        doc,
        "Подключена раздача статических файлов (public/css/style.css) для оформления интерфейса.",
    )

    add_subsection(doc, "2.2", "Шаблоны EJS")
    add_body(
        doc,
        "Создана папка views. Общий каркас вынесен в layout.ejs с областью <%- body %>. "
        "Страницы контента размещены в views/pages/: index.ejs (список кампаний), "
        "item.ejs (детальная информация по ID), add.ejs (форма добавления), "
        "login.ejs и register.ejs (вход и регистрация), а также страницы ошибок.",
    )
    add_paragraph(doc, "Рисунок 1 – Главная страница со списком кампаний", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    add_paragraph(doc, "(вставить скриншот http://localhost:3000/)", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    add_paragraph(doc, "Рисунок 2 – Форма добавления кампании", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    add_paragraph(doc, "(вставить скриншот /add после авторизации)", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)

    add_subsection(doc, "2.3", "Middleware")
    add_body(
        doc,
        "Логирующее middleware (middleware/requestLogger.js) выводит в консоль метку времени, "
        "HTTP-метод и URL каждого запроса; те же сведения записываются в файл logs/access.log "
        "через модуль utils/logger.js.",
    )
    add_body(
        doc,
        "Middleware авторизации (middleware/auth.js) совместимо с лабораторной работой №3: "
        "регистрация и вход через POST /auth/register и POST /auth/login (bcrypt + JWT). "
        "Для web-интерфейса JWT сохраняется в httpOnly-cookie; маршруты /add доступны "
        "только авторизованному пользователю, иначе выполняется redirect на /login.",
    )

    add_subsection(doc, "2.4", "Маршруты и данные")
    add_body(
        doc,
        "Маршруты web-интерфейса реализованы в routes/webRoutes.js, логика — в controllers/webController.js. "
        "GET / — список кампаний из models/campaignModel.js (массив в памяти); "
        "GET /item/:id — детальная страница; GET/POST /add — форма и добавление записи с редиректом на главную. "
        "REST API /campaigns сохранён для последующей интеграции с базой данных.",
    )

    add_subsection(doc, "2.5", "Обработка ошибок")
    add_body(
        doc,
        "В middleware/errorHandlers.js настроены обработчики 404 и 500: для web-запросов "
        "рендерятся шаблоны pages/not-found.ejs и pages/server-error.ejs, ошибки фиксируются в logs/error.log.",
    )
    add_paragraph(doc, "Рисунок 3 – Страница 404", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    add_paragraph(doc, "(вставить скриншот несуществующего URL)", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)

    add_section(doc, 3, "Контрольные вопросы")
    qa = [
        (
            "Что такое EJS и для чего он используется в Express?",
            "EJS (Embedded JavaScript) — шаблонизатор, позволяющий формировать HTML на сервере, "
            "подставляя данные из контроллера в разметку. Express подключает EJS через app.set('view engine', 'ejs') "
            "и метод res.render().",
        ),
        (
            "Как программно получить список элементов из in-memory хранилища?",
            "Данные кампаний хранятся в массиве в models/campaignModel.js; для чтения используются функции "
            "getAll() и getById(id), вызываемые из webController перед res.render().",
        ),
        (
            "Как организовано логирование запросов в файл?",
            "Middleware requestLogger вызывает logger.logRequest(), который дописывает строку "
            "с временем, методом и URL в logs/access.log (дополнительно — JSON-запись).",
        ),
        (
            "Объясните использование middleware в вашем приложении.",
            "Middleware — функции конвейера Express: requestLogger (лог), loadUserFromCookie (JWT из cookie), "
            "requireAuth (защита /add), notFoundHandler и errorHandler (404/500). "
            "Каждая функция получает (req, res, next) и передаёт управление дальше через next().",
        ),
    ]
    for i, (q, a) in enumerate(qa, 1):
        add_paragraph(doc, f"3.{i} {q}", bold=True, indent=True)
        add_body(doc, a)

    add_section(doc, 4, "Вывод")
    add_body(
        doc,
        "В ходе практической работы создано серверное web-приложение на Express и EJS "
        "в тематике курсового проекта. Реализованы логирование запросов в файл, "
        "регистрация и авторизация пользователей (JWT), динамические страницы списка, "
        "детализации и добавления кампаний, обработка ошибок 404 и 500. "
        "Структура auth-модулей согласована с лабораторной работой №3 для последующего объединения с БД.",
    )

    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build_document()
