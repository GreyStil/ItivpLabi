# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_PATH = Path(r"C:\Users\ASUS\Desktop\курс4\Lab22_Otchet_PostgreSQL_Sequelize.docx")


def set_run_font(run, name="Times New Roman", size=14, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_centered(doc, text, size=14, bold=False, space_after=0, space_before=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def add_body(doc, text, first_line_indent=True, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if first_line_indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    set_run_font(run, bold=bold)


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    set_run_font(run, bold=True)


def add_subheading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, bold=True)


def add_code(doc, text):
    for line in text.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(line if line else " ")
        set_run_font(run, name="Consolas", size=10)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=12)


def build():
    doc = Document()
    s = doc.sections[0]
    s.top_margin = Cm(2)
    s.bottom_margin = Cm(2)
    s.left_margin = Cm(3)
    s.right_margin = Cm(1.5)

    add_centered(doc, "Министерство образования Республики Беларусь")
    add_centered(doc, "Учреждение образования", space_before=6)
    add_centered(doc, "Белорусский государственный университет", bold=True)
    add_centered(doc, "информатики и радиоэлектроники", bold=True, space_after=12)
    add_centered(doc, "Факультет информационных технологий и управления")
    add_centered(doc, "Кафедра информационных технологий автоматизированных систем", space_after=18)
    add_centered(doc, "Дисциплина: Интернет-технологии и веб-программирование", space_after=36)
    add_centered(doc, "Отчёт", size=16, bold=True, space_before=24)
    add_centered(doc, "к лабораторной работе №2.2", bold=True)
    add_centered(doc, "PostgreSQL и Sequelize ORM", space_after=12)
    add_centered(doc, "Тема: REST API рекламных кампаний с постоянным хранением в БД", space_after=36)
    add_centered(doc, "Выполнил", space_before=24)
    add_centered(doc, "Р. С. Рубацкий", space_after=12)
    add_centered(doc, "Проверила")
    add_centered(doc, "Н. В. Хаджинова", space_after=48)
    for _ in range(4):
        doc.add_paragraph()
    add_centered(doc, "Минск 2026")
    doc.add_page_break()

    add_heading(doc, "1 ЦЕЛЬ РАБОТЫ")
    add_body(
        doc,
        "Заменить временное хранение данных в массиве (лабораторная работа №2.1) на "
        "постоянное хранение в PostgreSQL с использованием ORM Sequelize. Освоить настройку "
        "подключения к облачной БД Neon, миграции, сиды и CRUD-операции через методы Sequelize.",
    )

    add_heading(doc, "2 КРАТКИЕ ТЕОРЕТИЧЕСКИЕ СВЕДЕНИЯ ОБ ORM")
    add_body(
        doc,
        "ORM (Object-Relational Mapping) — технология, связывающая объекты приложения с "
        "таблицами реляционной базы данных. Разработчик работает с моделями и методами "
        "(findAll, create, update, destroy), а ORM формирует SQL-запросы автоматически.",
    )
    add_body(
        doc,
        "Sequelize — популярная ORM для Node.js. Она поддерживает миграции (версионирование "
        "схемы БД), seed-файлы (начальное наполнение), валидацию и транзакции. Это упрощает "
        "сопровождение проекта и интеграцию серверной части курсового приложения с PostgreSQL.",
    )

    add_heading(doc, "3 ХОД РАБОТЫ")

    add_subheading(doc, "3.1 Подготовка базы данных")
    add_body(
        doc,
        "Создан облачный проект в Neon (PostgreSQL). Проект связан с локальной папкой командой "
        "neon link, строка подключения сохранена в .env.local (переменная DATABASE_URL).",
    )
    add_caption(doc, "Рисунок 3.1 – Дашборд Neon / структура БД (вставить скриншот)")

    add_subheading(doc, "3.2 Установка зависимостей и инициализация Sequelize")
    add_body(doc, "Установлены пакеты: sequelize, pg, pg-hstore, dotenv, sequelize-cli.")
    add_code(doc, "npm install sequelize pg pg-hstore dotenv\nnpm install --save-dev sequelize-cli\nnpx sequelize-cli init")
    add_caption(doc, "Рисунок 3.2 – Конфигурация Sequelize (config/config.js, .sequelizerc)")

    add_subheading(doc, "3.3 Модель Campaign, миграции и сиды")
    add_body(
        doc,
        "Создана модель Campaign (аналог ресурса items из методички, но в предметной области "
        "рекламных кампаний). Выполнены миграции create-campaign и add-priority-to-campaigns "
        "(добавлено поле priority). Seed demo-campaigns заполнил таблицу двумя демонстрационными записями.",
    )
    add_code(
        doc,
        '''// migrations/...-add-priority-to-campaigns.js
await queryInterface.addColumn('Campaigns', 'priority', {
  type: Sequelize.INTEGER,
  allowNull: false,
  defaultValue: 1
});''',
    )
    add_caption(doc, "Рисунок 3.3 – Таблица Campaigns с полем priority (скриншот SQL Editor / Neon)")

    add_subheading(doc, "3.4 CRUD в Express через Sequelize")
    add_body(
        doc,
        "Маршруты /campaigns из ЛР №2.1 сохранены. Логика массива заменена на вызовы Sequelize "
        "в services/campaignService.js и controllers/campaignController.js:",
    )
    add_code(
        doc,
        '''async function getAll() {
  const rows = await Campaign.findAll({ order: [['id', 'ASC']] });
  return rows.map(formatCampaign);
}

async function getById(id) {
  const row = await Campaign.findByPk(id);
  return row ? formatCampaign(row) : null;
}

async function add(item) {
  const row = await Campaign.create({ ...item, ...calculateMetrics(item) });
  return formatCampaign(row);
}

async function remove(id) {
  const deleted = await Campaign.destroy({ where: { id } });
  return deleted > 0;
}''',
    )
    add_caption(doc, "Листинг 3.1 – Ключевые фрагменты CRUD (Sequelize)")

    add_subheading(doc, "3.5 Тестирование через Postman")
    add_body(
        doc,
        "После npm run db:migrate и npm run db:seed:all сервер запускается командой npm run dev. "
        "Проверены GET/POST/PUT/DELETE для /campaigns. Данные возвращаются из PostgreSQL, "
        "в том числе поле priority.",
    )
    add_caption(doc, "Рисунок 3.4 – GET /campaigns (Postman)")
    add_caption(doc, "Рисунок 3.5 – POST /campaigns (Postman)")
    add_caption(doc, "Рисунок 3.6 – PUT /campaigns/:id (Postman)")
    add_caption(doc, "Рисунок 3.7 – DELETE /campaigns/:id (Postman)")

    add_heading(doc, "4 ТАБЛИЦА МАРШРУТОВ")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Маршрут", "Sequelize-метод", "Статус", "Примечание"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h

    rows = [
        ("GET /campaigns", "Campaign.findAll()", "200", "Список из БД"),
        ("GET /campaigns/:id", "Campaign.findByPk(id)", "200/404", "Одна запись"),
        ("POST /campaigns", "Campaign.create(body)", "201/400", "Создание"),
        ("PUT /campaigns/:id", "instance.update(body)", "200/404", "Обновление"),
        ("DELETE /campaigns/:id", "Campaign.destroy({ where: { id } })", "204/404", "Удаление"),
    ]
    for row_data in rows:
        row = table.add_row()
        for i, val in enumerate(row_data):
            row.cells[i].text = val
    add_caption(doc, "Таблица 4.1 – Соответствие маршрутов методам Sequelize")

    add_heading(doc, "5 ОТВЕТЫ НА КОНТРОЛЬНЫЕ ВОПРОСЫ")
    qas = [
        (
            "1. Что такое ORM и зачем она нужна в Node.js-приложениях?",
            "ORM связывает объекты JavaScript с таблицами SQL-базы. Это снижает объём ручного SQL, "
            "ускоряет разработку CRUD и делает код проще для сопровождения.",
        ),
        (
            "2. Чем миграция отличается от seed?",
            "Миграция изменяет структуру БД (создание таблиц, добавление колонок). "
            "Seed заполняет таблицы начальными данными для разработки и тестирования.",
        ),
        (
            "3. Как Sequelize подключается к облачной PostgreSQL (Neon)?",
            "Через переменную окружения DATABASE_URL в config/config.js (dialect: postgres) "
            "и SSL-параметры для безопасного соединения.",
        ),
        (
            "4. Какие методы Sequelize использованы для CRUD?",
            "findAll, findByPk, create, update (через экземпляр модели), destroy.",
        ),
        (
            "5. Какие преимущества даёт переход с массива в памяти на PostgreSQL?",
            "Данные сохраняются после перезапуска сервера, возможна совместная работа нескольких "
            "клиентов, масштабирование и подключение полноценной СУБД в курсовом проекте.",
        ),
    ]
    for q, a in qas:
        add_body(doc, q, bold=True, first_line_indent=False)
        add_body(doc, a)

    add_heading(doc, "6 ВЫВОДЫ")
    add_body(
        doc,
        "Выполнена интеграция REST API рекламных кампаний с PostgreSQL через Sequelize. "
        "Настроены миграции, добавлено поле priority, выполнено начальное наполнение seed-данными. "
        "CRUD-маршруты из лабораторной работы №2.1 сохранены, изменена только логика хранения данных.",
    )
    add_body(
        doc,
        "Репозиторий: https://github.com/GreyStil/https---github.com-GreyStil-Itivp.git "
        "(ветка с выполнением лабораторной).",
    )

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_PATH))
    print(OUT_PATH)


if __name__ == "__main__":
    build()
